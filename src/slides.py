"""Render weekly update slides (PNG, 1920x1440) in the week-1/2 visual style, plus a .pptx wrapper.

A slide is a list of blocks stacked top to bottom:
  ('header', text) ('title', text) ('heading', text) ('text', text, {size, color, weight, after})
  ('table', columns, rows, widths) ('columns', left_blocks, right_blocks, split)
  ('hbar', {...}) ('line', {...}) ('heatmap', {...}) ('gap', px)
Numbers that appear on a slide should come from Registry.num(), so check_numbers() can prove every
number shown traces back to results.json.
"""
from __future__ import annotations
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from matplotlib.colors import TwoSlopeNorm  # noqa: E402

W, H, MARGIN = 1920, 1440, 96
FONT = ['Arial', 'DejaVu Sans']  # Helvetica Neue ships as a single-face .ttc here, so bold would silently drop
C = {'bg': '#F6F8FB', 'header': '#3C8DEB', 'title': '#141B2D', 'body': '#1F2937', 'muted': '#5B6472',
     'teal': '#117A7C', 'blue': '#3C8CFF', 'gray': '#A9B4C4', 'line': '#1F2937', 'target': '#D0473E',
     'white': '#FFFFFF', 'green': '#2E9E6A'}
PX = 100/72  # pixels per point at dpi 100


class Registry:
    """Formats numbers for slides and remembers every string it produced."""

    def __init__(self):
        self.shown = {}

    def num(self, value, fmt='{:.1f}', source=''):
        text = fmt.format(value)
        if text.startswith('-'):
            text = '−'+text[1:]  # typographic minus
        self.shown[text] = source
        return text


NUMBER = re.compile(r'\d+(?:,\d{3})*(?:\.\d+)?')


def check_numbers(texts, registry, allowed=()):
    """Numbers in the texts that were not produced by the registry (or explicitly allowed)."""
    known = {token for s in [*registry.shown, *allowed] for token in NUMBER.findall(s)}
    return sorted({token for text in texts for token in NUMBER.findall(text)}-known)


class _Canvas:
    def __init__(self):
        plt.rcParams['font.family'] = FONT
        self.fig = plt.figure(figsize=(W/100, H/100), dpi=100, facecolor=C['bg'])
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set_xlim(0, W)
        self.ax.set_ylim(H, 0)
        self.ax.axis('off')
        self.renderer = self.fig.canvas.get_renderer()
        self.texts = []

    def width(self, text, size, weight):
        t = self.ax.text(0, 0, text, fontsize=size, fontweight=weight)
        w = t.get_window_extent(self.renderer).width
        t.remove()
        return w

    def wrap(self, text, size, weight, max_width):
        lines = []
        for paragraph in text.split('\n'):
            line = ''
            for word in paragraph.split(' '):
                trial = f'{line} {word}'.strip()
                if line and self.width(trial, size, weight) > max_width:
                    lines.append(line)
                    line = word
                else:
                    line = trial
            lines.append(line)
        return lines

    def text(self, x, y, text, size=24, color=None, weight='normal', max_width=W-2*MARGIN, spacing=1.32):
        lines = self.wrap(text, size, weight, max_width)
        step = size*PX*spacing
        for i, line in enumerate(lines):
            self.ax.text(x, y+i*step, line, fontsize=size, color=color or C['body'], fontweight=weight, va='top')
        self.texts.append(text)
        return len(lines)*step

    def inset(self, x, y, w, h):
        return self.fig.add_axes([x/W, 1-(y+h)/H, w/W, h/H])


def _table(cv, x, y, width, columns, rows, widths, size=22):
    pad, top = 22, y
    edges = np.concatenate([[0], np.cumsum(widths)])*width
    for r, row in enumerate([columns]+rows):
        header = r == 0
        weight = 'bold' if header else 'normal'
        heights = [len(cv.wrap(cell, size, weight, (edges[i+1]-edges[i])-2*pad))*size*PX*1.3 for i, cell in enumerate(row)]
        h = max(heights)+2*pad
        cv.ax.add_patch(plt.Rectangle((x, y), width, h, facecolor=C['blue'] if header else C['white'],
                                      edgecolor=C['line'], linewidth=1.4))
        for i, cell in enumerate(row):
            if i:
                cv.ax.plot([x+edges[i]]*2, [y, y+h], color=C['line'], linewidth=1.4)
            cv.text(x+edges[i]+pad, y+pad, cell, size, C['white'] if header else C['body'], weight,
                    max_width=(edges[i+1]-edges[i])-2*pad, spacing=1.3)
        y += h
    return y-top


def _hbar(cv, x, y, width, spec):
    labels, values, colors = spec['labels'], spec['values'], spec['colors']
    h = spec.get('height', 60*len(labels)+90)
    label_w = spec.get('label_width', 300)
    ax = cv.inset(x+label_w, y, width-label_w-110, h-48)  # leave room for the x tick labels
    pos = np.arange(len(labels))[::-1]
    ax.barh(pos, values, color=colors, height=.55)
    ax.set_xlim(0, spec['xmax'])
    ax.set_ylim(-.6, len(labels)-.4+.7)  # headroom for the target label
    ax.set_yticks(pos, labels, fontsize=spec.get('label_size', 20), color=C['muted'])
    ax.tick_params(axis='x', labelsize=18, colors=C['body'], length=0)
    ax.tick_params(axis='y', length=0)
    for side in ['top', 'right', 'left']:
        ax.spines[side].set_visible(False)
    ax.spines['bottom'].set_color('#D5DBE3')
    ax.grid(axis='x', color='#E3E7ED', linewidth=1)
    ax.set_axisbelow(True)
    ax.set_facecolor(C['bg'])
    for p, v, text in zip(pos, values, spec['value_labels']):
        ax.text(min(v, spec['xmax'])+spec['xmax']*.015, p, text, va='center', fontsize=20, fontweight='bold', color=C['body'])
    if spec.get('target') is not None:
        ax.axvline(spec['target'], color=C['target'], linewidth=2.5, linestyle='--')
        ax.text(spec['target']+spec['xmax']*.012, len(labels)-.4+.35, spec['target_label'], color=C['target'], fontsize=17,
                fontweight='bold', ha='left', va='center')
    cv.texts.extend(spec['value_labels']+labels+[spec.get('target_label', '')])
    return h


def _line(cv, x, y, width, spec):
    h = spec.get('height', 380)
    ax = cv.inset(x+90, y, width-300, h-70)
    for series in spec['series']:
        ax.plot(spec['x'], series['y'], marker='o', linewidth=3, markersize=8, color=series['color'], label=series['label'])
    ax.legend(loc='upper right', fontsize=17, frameon=False)
    ax.set_xscale('log')
    ax.invert_xaxis()
    ax.set_xticks(spec['x'], spec['x_labels'], fontsize=17, color=C['body'])
    ax.minorticks_off()
    ax.tick_params(axis='y', labelsize=17, colors=C['body'], length=0)
    ax.tick_params(axis='x', length=0)
    ax.set_xlabel(spec['x_title'], fontsize=17, color=C['muted'])
    ax.set_ylabel(spec['y_title'], fontsize=17, color=C['muted'])
    for side in ['top', 'right']:
        ax.spines[side].set_visible(False)
    ax.grid(color='#E3E7ED', linewidth=1)
    ax.set_facecolor(C['bg'])
    if spec.get('target') is not None:
        ax.axhline(spec['target'], color=C['target'], linewidth=2.5, linestyle='--')
        ax.text(spec['x'][0], spec['target'], spec['target_label'], color=C['target'], fontsize=16, fontweight='bold',
                va='bottom', ha='left')
    cv.texts.extend(spec['x_labels']+[s['label'] for s in spec['series']]+[spec['x_title'], spec['y_title'], spec.get('target_label', '')])
    return h


def _heatmap(cv, x, y, width, spec):
    M = np.asarray(spec['matrix'], float)
    h = spec.get('height', 44*M.shape[0]+90)
    label_w, top_pad = spec.get('label_width', 150), spec.get('top_pad', 110)  # top_pad: room for rotated column names
    ax = cv.inset(x+label_w, y+top_pad, width-label_w-60, h-top_pad-10)
    lim = max(np.abs(M).max(), 1e-9)
    ax.imshow(M, cmap='RdBu_r', norm=TwoSlopeNorm(0, -lim, lim), aspect='auto')
    ax.set_yticks(range(M.shape[0]), spec['rows'], fontsize=15, color=C['body'])
    ax.xaxis.tick_top()
    ax.set_xticks(range(M.shape[1]), spec['cols'], fontsize=15, color=C['body'], rotation=spec.get('col_rotation', 35),
                  ha='left', rotation_mode='anchor')
    ax.tick_params(length=0)
    for side in ax.spines.values():
        side.set_visible(False)
    for (i, j), label in spec.get('annotations', {}).items():
        ax.text(j, i, label, ha='center', va='center', fontsize=14, fontweight='bold', color=C['white'])
    cv.texts.extend(spec.get('annotations', {}).values())  # row/column names are identifiers, not results
    return h


def _blocks(cv, blocks, x, y, width):
    top = y
    for block in blocks:
        kind = block[0]
        if kind == 'header':
            y += cv.text(x, y, block[1], 19, C['header'], 'bold', width)+26
        elif kind == 'title':
            y += cv.text(x, y, block[1], 36, C['title'], 'bold', width, spacing=1.2)+30
        elif kind == 'heading':
            y += cv.text(x, y, block[1], 23, C['teal'], 'bold', width)+14
        elif kind == 'text':
            opts = block[2] if len(block) > 2 else {}
            y += cv.text(x, y, block[1], opts.get('size', 23), opts.get('color', C['body']), opts.get('weight', 'normal'), width)+opts.get('after', 18)
        elif kind == 'table':
            y += _table(cv, x, y, width, block[1], block[2], block[3])+34
        elif kind == 'hbar':
            y += _hbar(cv, x, y, width, block[1])+18
        elif kind == 'heatmap':
            y += _heatmap(cv, x, y, width, block[1])+10
        elif kind == 'line':
            y += _line(cv, x, y, width, block[1])+10
        elif kind == 'gap':
            y += block[1]
        elif kind == 'columns':
            left, right, split = block[1], block[2], block[3]
            lw = (width-60)*split
            y += max(_blocks(cv, left, x, y, lw), _blocks(cv, right, x+lw+60, y, width-lw-60))+20
    return y-top


def render(blocks, path):
    """Draw one slide; returns (all text on the slide, bottom y in px) so callers can check fit and numbers."""
    cv = _Canvas()
    bottom = MARGIN*.6+_blocks(cv, blocks, MARGIN, MARGIN*.6, W-2*MARGIN)
    cv.fig.savefig(path, dpi=100, facecolor=C['bg'])
    plt.close(cv.fig)
    return cv.texts, bottom


def to_pptx(pngs, path):
    """Wrap rendered slides in a 4:3 deck (one full-bleed image per slide)."""
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    for png in pngs:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(png), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(path)
